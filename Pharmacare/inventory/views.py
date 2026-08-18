from django.shortcuts import render

# Create your views here.


import json
from datetime import date, timedelta
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.contrib.auth.decorators import login_required
from .models import Medicine
import openpyxl
from openpyxl.styles import PatternFill, Font
from io import BytesIO
from django.http import HttpResponse
from django.utils import timezone

def auth_required(request):
    if not request.user.is_authenticated:
        return JsonResponse({'error': 'Unauthorized'}, status=401)
    return None


@csrf_exempt
def medicines(request):
    auth = auth_required(request)
    if auth:
        return auth

    user = request.user

    if request.method == 'GET':
        meds = Medicine.objects.filter(user=user, exp_date__gte=date.today())
        data = []
        for m in meds:
            data.append({
                '_id': m.id,
                'name': m.name,
                'type': m.type,
                'composition': m.composition,
                'pack': m.pack,
                'place': m.place,
                'mfgDate': m.mfg_date,
                'expDate': m.exp_date,
                'mrp': m.mrp,
                'buyPrice': m.buy_price,
                'sellPrice': m.sell_price,
                'stock': m.stock,
                'lastUsed': m.last_used,
                'used_at': m.used_at,
                'createdAt': m.created_at
            })
        return JsonResponse({'medicines': data})

    if request.method == 'POST':
        data = json.loads(request.body)

        # Validate stock is not negative
        if data.get('stock', 0) < 0:
            return JsonResponse({'error': 'Stock cannot be negative'}, status=400)

        # Check for duplicate medicine with same name, composition, and expiry date
        # Allow duplicate if expiry date is different
        existing_med = Medicine.objects.filter(
            user=user,
            name=data['name'],
            composition=data.get('composition', ''),
            exp_date=data['expDate']
        ).first()

        if existing_med:
            return JsonResponse({
                'error': f"Medicine '{data['name']}' with expiry date {data['expDate']} already exists! Different expiry date allowed."
            }, status=400)

        med = Medicine.objects.create(
            user=user,
            name=data['name'],
            type=data.get('type', ''),
            composition=data.get('composition', ''),
            pack=data.get('pack', ''),
            place=data.get('place', ''),
            mfg_date=data.get('mfgDate') or None,
            exp_date=data['expDate'],
            mrp=data['mrp'],
            buy_price=data['buyPrice'],
            sell_price=data['sellPrice'],
            stock=data['stock'],
            used_at=timezone.now()
        )
        return JsonResponse({'message': 'Medicine added successfully'})

    return JsonResponse({'error': 'Method not allowed'}, status=405)


@csrf_exempt
def medicine_detail(request, id):
    auth = auth_required(request)
    if auth:
        return auth

    try:
        med = Medicine.objects.get(id=id, user=request.user)
    except Medicine.DoesNotExist:
        return JsonResponse({'error': 'Medicine not found'}, status=404)

    if request.method == 'PUT':
        data = json.loads(request.body)

        # Update used_at to current time for sorting by last used
        med.used_at = timezone.now()

        # Validate stock if provided
        if 'stock' in data and int(data.get('stock', 0)) < 0:
            return JsonResponse({'error': 'Stock cannot be negative'}, status=400)

        # If name/composition/expiry are provided, ensure we don't create a duplicate entry
        if 'name' in data and 'expDate' in data:
            existing_med = Medicine.objects.filter(
                user=request.user,
                name=data['name'],
                composition=data.get('composition', med.composition),
                exp_date=data['expDate']
            ).exclude(id=id).first()

            if existing_med:
                return JsonResponse({
                    'error': f"Medicine '{data['name']}' with expiry date {data['expDate']} already exists! Different expiry date allowed."
                }, status=400)

        if 'name' in data:
            med.name = data['name']
        if 'type' in data:
            med.type = data.get('type', '')
        if 'composition' in data:
            med.composition = data.get('composition', '')
        if 'pack' in data:
            med.pack = data.get('pack', '')
        if 'place' in data:
            med.place = data.get('place', '')
        if 'mfgDate' in data:
            med.mfg_date = data.get('mfgDate') or None
        if 'expDate' in data:
            med.exp_date = data['expDate']
        if 'mrp' in data:
            med.mrp = data['mrp']
        if 'buyPrice' in data:
            med.buy_price = data['buyPrice']
        if 'sellPrice' in data:
            med.sell_price = data['sellPrice']
        if 'stock' in data:
            med.stock = data['stock']
        med.save()
        return JsonResponse({'message': 'Medicine updated successfully'})

    if request.method == 'DELETE':
        med.delete()
        return JsonResponse({'message': 'Medicine deleted successfully'})

    return JsonResponse({'error': 'Method not allowed'}, status=405)



# Soon-Expiring medicines APIs

# Template view for soon-expiring page - renders HTML template
@login_required(login_url='/login.html')
def soon_expiring_page(request):
    """Render the soon-expiring medicines page"""
    return render(request, 'soon-expiring.html')


@csrf_exempt
def soon_expiring(request):
    auth = auth_required(request)
    if auth:
        return auth

    today = date.today()
    limit = today + timedelta(days=150)  # ~5 months

    meds = Medicine.objects.filter(
        user=request.user,
        exp_date__lte=limit
    )

    data = []
    for m in meds:
        data.append({
            '_id': m.id,
            'name': m.name,
            'type': m.type,
            'composition': m.composition,
            'pack': m.pack,
            'mfgDate': m.mfg_date,
            'expDate': m.exp_date,
            'mrp': m.mrp,
            'buyPrice': m.buy_price,
            'sellPrice': m.sell_price,
            'stock': m.stock,
            'lastUsed': m.last_used
        })

    return JsonResponse({'medicines': data})


# Expired medicines APIs

# Template view for expired medicines page - renders HTML template
@login_required(login_url='/login.html')
def expired_medicines_page(request):
    """Render the expired medicines page"""
    return render(request, 'expired-medicine.html')


@csrf_exempt
def expired_medicines(request):
    auth = auth_required(request)
    if auth:
        return auth

    today = date.today()

    # Get only expired medicines (exp_date < today)
    meds = Medicine.objects.filter(
        user=request.user,
        exp_date__lt=today
    ).order_by('exp_date')  # Oldest expired first

    data = []
    total_stock = 0
    total_value = 0
    
    for m in meds:
        stock_value = m.stock * m.mrp
        total_stock += m.stock
        total_value += stock_value
        
        data.append({
            '_id': m.id,
            'name': m.name,
            'type': m.type,
            'composition': m.composition,
            'pack': m.pack,
            'mfgDate': m.mfg_date,
            'expDate': m.exp_date,
            'mrp': m.mrp,
            'buyPrice': m.buy_price,
            'sellPrice': m.sell_price,
            'stock': m.stock,
            'stockValue': stock_value,
            'lastUsed': m.last_used,
            'daysExpired': (today - m.exp_date).days
        })

    return JsonResponse({
        'medicines': data,
        'totalExpired': len(data),
        'totalStock': total_stock,
        'totalValue': total_value
    })


@csrf_exempt
def export_medicines(request):
    """Export medicines to Excel file"""
    auth = auth_required(request)
    if auth:
        return auth
    
    try:
        meds = Medicine.objects.filter(user=request.user)
        
        # Create workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Medicines"
        
        # Add headers
        headers = ["ID", "Medicine Name", "Type", "Composition", "MFG Date", "EXP Date", "MRP", "Buy Price", "Sell Price", "Stock"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")
        
        # Add data rows
        for row, m in enumerate(meds, 2):
            ws.cell(row=row, column=1, value=m.id)
            ws.cell(row=row, column=2, value=m.name)
            ws.cell(row=row, column=3, value=m.type)
            ws.cell(row=row, column=4, value=m.composition)
            ws.cell(row=row, column=3, value=m.pack)
            ws.cell(row=row, column=5, value=m.mfg_date.strftime('%Y-%m-%d') if m.mfg_date else '')
            ws.cell(row=row, column=6, value=m.exp_date.strftime('%Y-%m-%d') if m.exp_date else '')
            ws.cell(row=row, column=7, value=m.mrp)
            ws.cell(row=row, column=8, value=m.buy_price)
            ws.cell(row=row, column=9, value=m.sell_price)
            ws.cell(row=row, column=10, value=m.stock)
        
        # Adjust column widths
        for col in range(1, 11):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18
        
        # Write to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="medicines_{date.today()}.xlsx"'
        return response
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)


@csrf_exempt
def import_medicines(request):
    """Bulk import medicines from Excel file or JSON data"""
    auth = auth_required(request)
    if auth:
        return auth

    if request.method != 'POST':
        return JsonResponse({'error': 'Method not allowed'}, status=405)

    try:
        created_count = 0
        skipped_count = 0
        medicines_data = []

        if 'file' in request.FILES:
            # Handle file upload
            excel_file = request.FILES['file']
            wb = openpyxl.load_workbook(excel_file)
            ws = wb.active

            # Expected headers: Medicine Name, Type, Composition, MFG Date, EXP Date, MRP, Buy Price, Sell Price, Stock (ID is optional and ignored)
            headers = [cell.value for cell in ws[1]]
            expected_headers = ["Medicine Name", "Type", "Composition", "MFG Date", "EXP Date", "MRP", "Buy Price", "Sell Price", "Stock"]
            # Normalize headers for comparison (case-insensitive, strip whitespace)
            normalized_headers = [str(h).strip().lower() if h is not None else '' for h in headers]
            normalized_expected = [eh.strip().lower() for eh in expected_headers]
            # Create header map for flexible order
            header_map = {norm_h: i for i, norm_h in enumerate(normalized_headers)}
            if not all(eh in header_map for eh in normalized_expected):
                return JsonResponse({'error': 'Invalid Excel format. Expected headers: ' + ', '.join(expected_headers)}, status=400)

            for row in ws.iter_rows(min_row=2, values_only=True):
                try:
                    row_values = list(row)
                    name = row_values[header_map['medicine name']]
                    med_type = row_values[header_map['type']]
                    composition = row_values[header_map['composition']]
                    pack = row_values[header_map['pack']] if 'pack' in header_map else ''
                    mfg_date_str = row_values[header_map['mfg date']]
                    exp_date_str = row_values[header_map['exp date']]
                    mrp = row_values[header_map['mrp']]
                    buy_price = row_values[header_map['buy price']]
                    sell_price = row_values[header_map['sell price']]
                    stock = row_values[header_map['stock']]

                    # Parse dates
                    mfg_date = None
                    if mfg_date_str:
                        try:
                            mfg_date = date.fromisoformat(str(mfg_date_str))
                        except:
                            mfg_date = None

                    exp_date = None
                    if exp_date_str:
                        try:
                            exp_date = date.fromisoformat(str(exp_date_str))
                        except:
                            skipped_count += 1
                            continue  # Skip if exp_date is invalid

                    # Convert prices and stock
                    try:
                        mrp = float(mrp) if mrp else 0
                    except:
                        mrp = 0

                    try:
                        buy_price = float(buy_price) if buy_price else 0
                    except:
                        buy_price = 0

                    try:
                        sell_price = float(sell_price) if sell_price else 0
                    except:
                        sell_price = 0

                    try:
                        stock = int(stock) if stock else 0
                    except:
                        stock = 0

                    # Prevent negative stock
                    if stock < 0:
                        skipped_count += 1
                        continue

                    medicines_data.append({
                        'name': name or '',
                        'type': med_type or '',
                        'composition': composition or '',
                        'pack': pack or '',
                        'mfg_date': mfg_date,
                        'exp_date': exp_date,
                        'mrp': mrp,
                        'buy_price': buy_price,
                        'sell_price': sell_price,
                        'stock': stock
                    })
                except Exception as e:
                    print(f"Error parsing row: {e}")
                    skipped_count += 1
                    continue
        else:
            # Handle JSON data
            data = json.loads(request.body)
            medicines_data = data.get('medicines', [])

        # Process medicines_data
        for med_data in medicines_data:
            try:
                name = med_data.get('name', '').strip()
                composition = med_data.get('composition', '').strip()
                exp_date = med_data.get('expDate') or med_data.get('exp_date')

                if not name or not exp_date:
                    skipped_count += 1
                    continue

                # Parse exp_date if string
                if isinstance(exp_date, str):
                    try:
                        exp_date = date.fromisoformat(exp_date)
                    except:
                        skipped_count += 1
                        continue

                # Check for duplicate
                existing_med = Medicine.objects.filter(
                    user=request.user,
                    name=name,
                    composition=composition,
                    exp_date=exp_date
                ).first()

                if existing_med:
                    skipped_count += 1
                    continue

                Medicine.objects.create(
                    user=request.user,
                    name=name,
                    type=med_data.get('type', ''),
                    composition=composition,
                    pack=med_data.get('pack', ''),
                    mfg_date=med_data.get('mfgDate') or med_data.get('mfg_date'),
                    exp_date=exp_date,
                    mrp=med_data.get('mrp', 0),
                    buy_price=med_data.get('buyPrice', 0) or med_data.get('buy_price', 0),
                    sell_price=med_data.get('sellPrice', 0) or med_data.get('sell_price', 0),
                    stock=med_data.get('stock', 0),
                    used_at=timezone.now()
                )
                created_count += 1
            except Exception as e:
                print(f"Error creating medicine: {e}")
                skipped_count += 1
                continue

        # After import, check for expired medicines
        expired_meds = Medicine.objects.filter(
            user=request.user,
            exp_date__lt=date.today()
        )

        expired_data = []
        for m in expired_meds:
            expired_data.append({
                '_id': m.id,
                'name': m.name,
                'type': m.type,
                'composition': m.composition,
                'pack': m.pack,
                'mfgDate': m.mfg_date,
                'expDate': m.exp_date,
                'mrp': m.mrp,
                'buyPrice': m.buy_price,
                'sellPrice': m.sell_price,
                'stock': m.stock,
                'lastUsed': m.last_used,
                'used_at': m.used_at,
                'createdAt': m.created_at
            })

        response_data = {
            'message': f'Successfully imported {created_count} medicines. Skipped {skipped_count} due to errors or invalid data.',
            'count': created_count,
            'skipped': skipped_count
        }

        if expired_data:
            response_data['expired_medicines'] = expired_data
            response_data['message'] += f' Found {len(expired_data)} expired medicines. Use edit/delete functions on them.'

        return JsonResponse(response_data)
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)


@csrf_exempt
def search_medicines(request):
    """Search medicines by name or composition"""
    auth = auth_required(request)
    if auth:
        return auth
    
    query = request.GET.get('q', '').strip()
    if not query:
        return JsonResponse({'medicines': []})
    
    try:
        from django.db.models import Q
        query_filter = Q(name__icontains=query) | Q(composition__icontains=query)
        meds = Medicine.objects.filter(
            query_filter,
            user=request.user
        )
        
        data = []
        for m in meds:
            data.append({
                '_id': m.id,
                'name': m.name,
                'type': m.type,
                'composition': m.composition,
                'pack': m.pack,
                'mfgDate': m.mfg_date,
                'expDate': m.exp_date,
                'mrp': m.mrp,
                'buyPrice': m.buy_price,
                'sellPrice': m.sell_price,
                'stock': m.stock,
                'lastUsed': m.last_used
            })
        
        return JsonResponse({'medicines': data})
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

from django.shortcuts import render, get_object_or_404
from .models import Medicine

def dashboard(request):
    medicine_id = request.GET.get('medicine_id')
    medicine = None

    if medicine_id:
        medicine = get_object_or_404(Medicine, id=medicine_id)

    return render(request, 'dashboard.html', {
        'medicine': medicine
    })

@login_required(login_url='/login.html')
def expired_medicines_page(request):
    """Render the expired medicines page"""
    return render(request, 'expired-medicine.html')